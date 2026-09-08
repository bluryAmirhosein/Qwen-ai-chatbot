import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.services.file_service import FileService, UnsupportedFileTypeError


class _FakeUploadFile:
    """Minimal stand-in for fastapi.UploadFile: only `.filename` and
    `.file` (a file-like object) are used by FileService.
    """

    def __init__(self, filename: str | None, content: bytes):
        self.filename = filename
        self.file = io.BytesIO(content)


@pytest.fixture
def service():
    return FileService()


def test_extract_text_from_txt(service):
    upload = _FakeUploadFile("notes.txt", "hello world".encode("utf-8"))
    assert service.extract_text(upload) == "hello world"


def test_extract_text_from_md(service):
    upload = _FakeUploadFile("notes.md", "# Title\n\nBody".encode("utf-8"))
    assert service.extract_text(upload) == "# Title\n\nBody"


def test_extract_text_ignores_undecodable_bytes_instead_of_raising(service):
    upload = _FakeUploadFile("notes.txt", b"valid \xff\xfe bytes")
    # Should not raise a UnicodeDecodeError.
    result = service.extract_text(upload)
    assert "valid" in result


def test_extract_text_is_case_insensitive_on_extension(service):
    upload = _FakeUploadFile("NOTES.TXT", b"hi")
    assert service.extract_text(upload) == "hi"


def test_extract_text_from_docx(service):
    document = Document()
    document.add_paragraph("First paragraph.")
    document.add_paragraph("Second paragraph.")
    buf = io.BytesIO()
    document.save(buf)

    upload = _FakeUploadFile("report.docx", buf.getvalue())
    result = service.extract_text(upload)

    assert result == "First paragraph.\nSecond paragraph."


def test_extract_text_from_pdf_does_not_raise_on_blank_page(service):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buf = io.BytesIO()
    writer.write(buf)

    upload = _FakeUploadFile("scan.pdf", buf.getvalue())
    result = service.extract_text(upload)

    assert result == ""


@pytest.mark.parametrize("filename", ["archive.zip", "image.png", "data.csv"])
def test_extract_text_rejects_unsupported_extensions(service, filename):
    upload = _FakeUploadFile(filename, b"irrelevant")
    with pytest.raises(UnsupportedFileTypeError):
        service.extract_text(upload)


def test_extract_text_rejects_missing_filename(service):
    upload = _FakeUploadFile(None, b"irrelevant")
    with pytest.raises(UnsupportedFileTypeError):
        service.extract_text(upload)


def test_extract_text_rejects_filename_without_extension(service):
    upload = _FakeUploadFile("README", b"irrelevant")
    with pytest.raises(UnsupportedFileTypeError):
        service.extract_text(upload)
