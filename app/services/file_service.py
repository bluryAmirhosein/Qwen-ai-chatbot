import io
import logging

from docx import Document
from fastapi import UploadFile
from openpyxl import load_workbook
from pypdf import PdfReader

logger = logging.getLogger(__name__)

_EXCEL_TYPES = {".xlsx", ".xlsm"}
_SUPPORTED_TYPES = {".txt", ".md", ".pdf", ".docx", ".csv"} | _EXCEL_TYPES

# Upper bound for the flat-text rendering of a workbook (used when an Excel
# file is attached directly to a chat message rather than ingested into RAG).
# A 1000x30 sheet easily exceeds what a small local model can take as context.
_EXCEL_TEXT_MAX_CHARS = 20_000


class UnsupportedFileTypeError(ValueError):
    pass


class FileService:
    """Extracts plain text from uploaded files so it can be fed to the
    model as context. Kept separate from ChatService so new file types
    can be added without touching conversation logic.

    Excel files have two paths:
    - `read_bytes` hands the raw workbook to ExcelChunker (RAG ingestion),
      which keeps sheet/row/column structure.
    - `extract_text` renders a truncated flat-text view, for callers that
      just need something to put in a prompt.
    """

    def extract_text(self, file: UploadFile) -> str:
        suffix = self._ensure_supported(file.filename)
        raw = file.file.read()

        if suffix in (".txt", ".md"):
            return raw.decode("utf-8", errors="ignore")

        if suffix == ".pdf":
            return self._extract_pdf(raw)

        if suffix == ".docx":
            return self._extract_docx(raw)

        if suffix == ".csv":
            return self._extract_csv(raw)

        if suffix in _EXCEL_TYPES:
            return self._extract_excel(raw)

        raise UnsupportedFileTypeError(f"Unsupported file type: {suffix}")

    def read_bytes(self, file: UploadFile) -> bytes:
        """Return the raw bytes of an uploaded file after validating its type."""
        self._ensure_supported(file.filename)
        return file.file.read()

    @staticmethod
    def is_csv(filename: str | None) -> bool:
        return FileService._get_suffix(filename) == ".csv"

    @staticmethod
    def is_excel(filename: str | None) -> bool:
        return FileService._get_suffix(filename) in _EXCEL_TYPES

    @staticmethod
    def _get_suffix(filename: str | None) -> str:
        if not filename or "." not in filename:
            return ""
        return "." + filename.rsplit(".", 1)[-1].lower()

    @staticmethod
    def _ensure_supported(filename: str | None) -> str:
        suffix = FileService._get_suffix(filename)
        if suffix not in _SUPPORTED_TYPES:
            supported = ", ".join(sorted(_SUPPORTED_TYPES))
            raise UnsupportedFileTypeError(
                f"Unsupported file type: '{suffix or 'unknown'}'. Supported types: {supported}"
            )
        return suffix

    @staticmethod
    def _extract_pdf(raw: bytes) -> str:
        reader = PdfReader(io.BytesIO(raw))
        return "\n".join(page.extract_text() or "" for page in reader.pages)

    @staticmethod
    def _extract_docx(raw: bytes) -> str:
        document = Document(io.BytesIO(raw))
        return "\n".join(p.text for p in document.paragraphs)

    @staticmethod
    def _extract_csv(raw: bytes) -> str:
        # utf-8-sig strips a leading BOM, which is common in CSVs exported
        # from Excel. cp1256 is the legacy Windows Arabic/Persian code page;
        # CSVs saved by older Excel versions in Iran are often in it, and
        # decoding them as latin-1 would turn all Persian text into garbage.
        # The result is handed off as-is to CsvChunker, which does the actual
        # row parsing.
        for encoding in ("utf-8-sig", "cp1256"):
            try:
                return raw.decode(encoding)
            except UnicodeDecodeError:
                logger.warning("CSV file was not valid %s; trying next encoding", encoding)
        return raw.decode("latin-1", errors="replace")

    @staticmethod
    def _extract_excel(raw: bytes) -> str:
        workbook = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
        try:
            parts: list[str] = []
            total = 0
            for sheet in workbook.worksheets:
                parts.append(f"## Sheet: {sheet.title}")
                for row in sheet.iter_rows(values_only=True):
                    cells = ["" if cell is None else str(cell).strip() for cell in row]
                    if not any(cells):
                        continue
                    line = " | ".join(cells)
                    total += len(line) + 1
                    if total > _EXCEL_TEXT_MAX_CHARS:
                        parts.append("[... truncated: workbook is too large to include in full ...]")
                        return "\n".join(parts)
                    parts.append(line)
            return "\n".join(parts)
        finally:
            workbook.close()